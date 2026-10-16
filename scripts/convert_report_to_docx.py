"""
Converts BTech_Project2_Report_DDU.md into a professional Microsoft Word (.docx) document
following Dharmsinh Desai University (DDU) format.
"""
import os
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    """Sets background color of a table cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets padding/margins for a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_styled_paragraph(doc, text, style='Normal', space_after=6, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = align
    
    # Simple markdown inline formatting: bold **text**, code `text`
    parts = re.split(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) >= 4:
            run = p.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*') and len(part) >= 2:
            run = p.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`') and len(part) >= 2:
            run = p.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(180, 40, 40)
        else:
            p.add_run(part)
    return p

def main():
    md_path = Path("docs/BTech_Project2_Report_DDU.md")
    if not md_path.exists():
        print(f"Error: {md_path} not found.")
        return

    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()

    # Set page margins (1 inch all sides)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Configure Normal Style Font
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 30, 30)

    in_table = False
    table_lines = []
    in_code_block = False
    code_lines = []

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\r\n')

        # Check for code block fences
        if line.startswith("```"):
            if in_code_block:
                # End of code block
                in_code_block = False
                code_text = "\n".join(code_lines)
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.left_indent = Inches(0.4)
                run = p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(40, 40, 40)
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Check for table
        if line.strip().startswith('|') and '|' in line.strip()[1:]:
            in_table = True
            table_lines.append(line.strip())
            i += 1
            continue
        elif in_table:
            # End of table, process it
            in_table = False
            # Parse table_lines
            raw_rows = []
            for t_line in table_lines:
                # Strip leading and trailing '|'
                content = t_line.strip()
                if content.startswith('|'):
                    content = content[1:]
                if content.endswith('|'):
                    content = content[:-1]
                cols = [c.strip() for c in content.split('|')]
                # Skip divider lines e.g. | :--- | :--- |
                if all(re.match(r'^:?-+:?$', c) for c in cols if c):
                    continue
                raw_rows.append(cols)

            if raw_rows:
                num_cols = max(len(r) for r in raw_rows)
                tbl = doc.add_table(rows=len(raw_rows), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                tbl.autofit = True
                
                # Apply Table Grid style
                tbl.style = 'Table Grid'

                for r_idx, row in enumerate(raw_rows):
                    for c_idx in range(num_cols):
                        cell_val = row[c_idx] if c_idx < len(row) else ""
                        cell = tbl.cell(r_idx, c_idx)
                        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
                        
                        cp = cell.paragraphs[0]
                        cp.paragraph_format.space_after = Pt(2)
                        cp.paragraph_format.space_before = Pt(2)
                        cp.paragraph_format.line_spacing = 1.05
                        
                        if r_idx == 0:
                            # Header row formatting
                            set_cell_background(cell, "E2E8F0")
                            run = cp.add_run(cell_val.replace("**", ""))
                            run.bold = True
                            run.font.name = 'Times New Roman'
                            run.font.size = Pt(10)
                            run.font.color.rgb = RGBColor(15, 23, 42)
                        else:
                            # Alternate row shading
                            if r_idx % 2 == 1:
                                set_cell_background(cell, "F8FAFC")
                            # Simple inline bold
                            parts = re.split(r'(\*\*.*?\*\*)', cell_val)
                            for part in parts:
                                if part.startswith('**') and part.endswith('**'):
                                    r = cp.add_run(part[2:-2])
                                    r.bold = True
                                else:
                                    r = cp.add_run(part)
                                r.font.name = 'Times New Roman'
                                r.font.size = Pt(9.5)
                doc.add_paragraph().paragraph_format.space_after = Pt(4)

            table_lines = []
            # do not increment i, handle current line below

        # Check for page break
        if line.strip() == r"\newpage" or line.strip() == "---":
            doc.add_page_break()
            i += 1
            continue

        # Check for Headings
        if line.startswith("# "):
            text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(18)
            run.font.color.rgb = RGBColor(15, 23, 42)
        elif line.startswith("## "):
            text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(30, 41, 59)
        elif line.startswith("### "):
            text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(51, 65, 85)
        elif line.startswith("#### "):
            text = line[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.bold = True
            run.italic = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(71, 85, 105)
        elif line.startswith("* ") or line.startswith("- "):
            text = line[2:].strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            # inline formatting
            parts = re.split(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                elif part.startswith('*') and part.endswith('*'):
                    run = p.add_run(part[1:-1])
                    run.italic = True
                elif part.startswith('`') and part.endswith('`'):
                    run = p.add_run(part[1:-1])
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9.5)
                else:
                    p.add_run(part)
        elif re.match(r'^\d+\.\s', line):
            # Numbered list
            idx_end = line.find(' ')
            text = line[idx_end+1:].strip()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            parts = re.split(r'(\*\*.*?\*\*|`.*?`|\*.*?\*)', text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.bold = True
                elif part.startswith('*') and part.endswith('*'):
                    run = p.add_run(part[1:-1])
                    run.italic = True
                elif part.startswith('`') and part.endswith('`'):
                    run = p.add_run(part[1:-1])
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9.5)
                else:
                    p.add_run(part)
        elif line.startswith("> "):
            # Blockquote / note
            text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(text)
            run.italic = True
            run.font.color.rgb = RGBColor(70, 70, 70)
        elif line.strip() == "":
            pass
        else:
            add_styled_paragraph(doc, line.strip())

        i += 1

    # Save to targets
    targets = [
        Path("docs/BTech_Project2_Report_DDU.docx"),
        Path("BTech_Project2_Report_DDU.docx"),
        Path(os.path.expanduser("~")) / "OneDrive" / "Desktop" / "BTech_Project2_Report_DDU.docx",
    ]

    for tgt in targets:
        try:
            tgt.parent.mkdir(parents=True, exist_ok=True)
            doc.save(str(tgt))
            print(f"[SUCCESS] Saved Word document: {tgt}")
        except Exception as e:
            print(f"[Warning] Could not write to {tgt}: {e}")

if __name__ == "__main__":
    main()
