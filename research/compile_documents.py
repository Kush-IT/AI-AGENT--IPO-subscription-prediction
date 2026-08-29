"""
Script to compile the academic research paper into publication-grade DOCX and PDF documents.
"""
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, KeepTogether, PageBreak, HRFlowable
)
from reportlab.pdfgen import canvas

BASE_DIR = Path(__file__).parent
PAPER_DIR = BASE_DIR / "paper"
FIGURES_DIR = BASE_DIR / "figures"
PAPER_DIR.mkdir(parents=True, exist_ok=True)

DOCX_PATH = PAPER_DIR / "IPO_Agentic_AI_Research_Paper.docx"
PDF_PATH = PAPER_DIR / "IPO_Agentic_AI_Research_Paper.pdf"


def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)


# =====================================================================
# 1. BUILD DOCX
# =====================================================================
def build_docx():
    print("Generating Word document (.docx)...")
    doc = docx.Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("Leakage-Aware Machine Learning and Agentic AI for Pre-IPO Subscription Demand Forecasting Across QIB, NII, Retail, and Total Subscription")
    r_title.bold = True
    r_title.font.size = Pt(18)
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    # Authors
    p_auth = doc.add_paragraph()
    p_auth.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_auth = p_auth.add_run("Quantitative Financial Engineering and Computational Intelligence Group\nDataset Version 1.1.0 | Model Registry Version 1.0.0-final-candidates")
    r_auth.font.size = Pt(10)
    r_auth.font.italic = True
    r_auth.font.color.rgb = RGBColor(71, 85, 105)

    doc.add_paragraph() # Spacer

    # Abstract Callout Box
    table_abs = doc.add_table(rows=1, cols=1)
    table_abs.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_abs = table_abs.cell(0, 0)
    set_cell_background(cell_abs, "F1F5F9")
    set_cell_margins(cell_abs, top=140, bottom=140, left=200, right=200)

    p_abs = cell_abs.paragraphs[0]
    r_abs_label = p_abs.add_run("Abstract— ")
    r_abs_label.bold = True
    r_abs_label.font.size = Pt(10)
    r_abs_text = p_abs.add_run(
        "Accurately forecasting subscription demand for Initial Public Offerings (IPOs) prior to bidding is a fundamental challenge in financial econometrics. "
        "In this paper, we propose an ex-ante, leakage-aware Machine Learning (ML) and Agentic Artificial Intelligence (AI) framework designed to forecast category-wise "
        "subscription demand multiples across Qualified Institutional Buyers (QIB), Non-Institutional Investors (NII/HNI), Retail Individual Investors (RII), and Total Aggregate demand for Indian Mainboard IPOs. "
        "We construct a curated dataset of 106 verified Indian Mainboard IPOs (2021–2024), governed by an inviolable ex-ante temporal boundary (t_observed <= 09:59:59 IST on open date). "
        "We formulate 45 predictive features across seven cumulative ablation tiers. On a quarantined 2024 out-of-sample test set of 17 verified Mainboard IPOs, the frozen Total model achieves R² = +0.666, MAE = 17.43x, and MedianAE = 14.84x, "
        "while the QIB model achieves R² = +0.589 and MAE = 31.43x. Incorporating pre-open GMP alongside media sentiment yields a 41.568% predictive error reduction over pure fundamental baselines. "
        "We encapsulate this system within an 8-node LangGraph autonomous agent orchestration graph, TreeSHAP explainability, and a React/TypeScript research dashboard with 100% test verification (98/98 tests passed)."
    )
    r_abs_text.font.size = Pt(9.5)

    p_kw = cell_abs.add_paragraph()
    r_kw_label = p_kw.add_run("\nKeywords: ")
    r_kw_label.bold = True
    r_kw_label.font.size = Pt(9)
    r_kw_text = p_kw.add_run("Initial Public Offerings (IPOs), Subscription Demand Forecasting, Temporal Data Leakage, Explainable AI (XAI), TreeSHAP, Agentic AI, LangGraph, Decision Support Systems, Emerging Markets.")
    r_kw_text.font.size = Pt(9)
    r_kw_text.font.italic = True

    doc.add_paragraph() # Spacer

    # Section 1: Introduction
    h1 = doc.add_heading("1. Introduction", level=1)
    doc.add_paragraph(
        "Initial Public Offerings (IPOs) represent a critical corporate transition wherein privately held enterprises issue equity shares to the public to raise long-term capital, optimize balance sheet leverage, and establish market valuation. Over the past decade, the Indian primary capital market has emerged as one of the most vibrant issuance ecosystems globally, leading worldwide rankings in the number of public offerings in recent years."
    )
    doc.add_paragraph(
        "Under the regulatory oversight of the Securities and Exchange Board of India (SEBI) and its Issue of Capital and Disclosure Requirements (ICDR) regulations, book-built IPOs allocate shares across legally defined investor categories: Qualified Institutional Buyers (QIB), Non-Institutional Investors (NII), and Retail Individual Investors (RII). Anticipating subscription demand across these tiers is paramount for institutional underwriters, retail market participants, corporate treasuries, and regulatory bodies."
    )

    # Insert Figure 1
    if (FIGURES_DIR / "fig01_methodology.png").exists():
        doc.add_picture(str(FIGURES_DIR / "fig01_methodology.png"), width=Inches(6.0))
        p_cap = doc.add_paragraph("Figure 1: Overall Pre-IPO Demand Forecasting & Explainability Research Pipeline.")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.runs[0].font.size = Pt(8.5)
        p_cap.runs[0].font.italic = True

    # Section 2: Temporal Anti-Leakage
    doc.add_heading("2. Strict Pre-IPO Temporal Anti-Leakage Framework", level=1)
    doc.add_paragraph(
        "To eliminate temporal look-ahead contamination, we establish a formal temporal validity condition. For any observed feature x_ij observed at timestamp t_observed and statutory opening timestamp t_cutoff (09:59:59 IST on opening date), the observation is valid if and only if t_observed <= t_cutoff. All post-open order book data, intraday bidding tallies, and listing returns are strictly rejected."
    )

    if (FIGURES_DIR / "fig02_temporal_leakage.png").exists():
        doc.add_picture(str(FIGURES_DIR / "fig02_temporal_leakage.png"), width=Inches(6.0))
        p_cap = doc.add_paragraph("Figure 2: Strict Pre-IPO Ex-Ante Temporal Anti-Leakage Boundary.")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.runs[0].font.size = Pt(8.5)
        p_cap.runs[0].font.italic = True

    # Section 3: Feature Engineering & Ablation
    doc.add_heading("3. Feature Engineering & Hierarchical Ablation Hierarchy", level=1)
    doc.add_paragraph(
        "We construct 45 predictive features across seven cumulative ablation tiers (Groups A through G). Group A captures issue terms, Group B incorporates financial fundamentals, Group C adds valuation ratios, Group D injects secondary market momentum, Group E adds historical past-only demand, Group F incorporates FinBERT sentiment and pre-open GMP, and Group G completes the matrix with sector encodings."
    )

    if (FIGURES_DIR / "fig03_feature_engineering.png").exists():
        doc.add_picture(str(FIGURES_DIR / "fig03_feature_engineering.png"), width=Inches(6.0))
        p_cap = doc.add_paragraph("Figure 3: Cumulative Hierarchical Feature Engineering Ablation Hierarchy (Groups A–G).")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.runs[0].font.size = Pt(8.5)
        p_cap.runs[0].font.italic = True

    # Section 4: Quantitative Results Table
    doc.add_heading("4. Quarantined Test Evaluation & Model Selection", level=1)
    doc.add_paragraph(
        "The frozen models were evaluated once on the strictly sealed 2024 test set (N=17 verified Mainboard IPOs). Table 1 presents the quantitative out-of-sample performance across all four target dimensions."
    )

    # Add Table
    table_res = doc.add_table(rows=5, cols=6)
    table_res.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Target Category", "Selected Model", "Target Transform", "Test MAE", "Test RMSE", "Test R²"]
    for j, h in enumerate(headers):
        cell = table_res.cell(0, j)
        cell.paragraphs[0].add_run(h).bold = True
        set_cell_background(cell, "0F172A")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        cell.paragraphs[0].runs[0].font.size = Pt(8.5)

    data_rows = [
        ["Total Subscription", "XGBoost Regressor", "log1p", "17.43x", "22.92x", "+0.666"],
        ["QIB Subscription", "Random Forest", "RAW", "31.43x", "38.48x", "+0.589"],
        ["NII Subscription", "LightGBM Regressor", "log1p", "39.62x", "64.44x", "+0.385"],
        ["Retail Subscription", "XGBoost Regressor", "RAW", "23.14x", "36.71x", "-0.450"],
    ]
    for i, row in enumerate(data_rows):
        for j, val in enumerate(row):
            cell = table_res.cell(i+1, j)
            cell.paragraphs[0].add_run(val)
            cell.paragraphs[0].runs[0].font.size = Pt(8.5)
            if i % 2 == 1:
                set_cell_background(cell, "F8FAFC")

    doc.add_paragraph() # Spacer

    if (FIGURES_DIR / "fig08_actual_vs_predicted.png").exists():
        doc.add_picture(str(FIGURES_DIR / "fig08_actual_vs_predicted.png"), width=Inches(4.5))
        p_cap = doc.add_paragraph("Figure 4: Actual vs. Predicted Total Subscription on 17 Quarantined 2024 Test IPOs (R² = +0.666).")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.runs[0].font.size = Pt(8.5)
        p_cap.runs[0].font.italic = True

    # Section 5: TreeSHAP Explainability
    doc.add_heading("5. TreeSHAP Additive Feature Explainability", level=1)
    doc.add_paragraph(
        "Predictions are decomposed via TreeSHAP into additive feature attributions relative to the historical base rate. Pre-open GMP, news sentiment, issue size friction, and market volatility emerge as dominant drivers across investor classes."
    )

    if (FIGURES_DIR / "fig07_shap_importance.png").exists():
        doc.add_picture(str(FIGURES_DIR / "fig07_shap_importance.png"), width=Inches(5.5))
        p_cap = doc.add_paragraph("Figure 5: Global TreeSHAP Feature Attributions for Total Subscription (XGBoost log1p).")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.runs[0].font.size = Pt(8.5)
        p_cap.runs[0].font.italic = True

    # Section 6: Agentic AI & UI
    doc.add_heading("6. Agentic AI Orchestration & Full-Stack Deployment", level=1)
    doc.add_paragraph(
        "We implement an 8-node LangGraph StateGraph pipeline coordinating input validation, data fetching, feature preparation, prediction, SHAP explanation, category ranking, and risk auditing into a responsive React/TypeScript dashboard with Docker containerization."
    )

    if (FIGURES_DIR / "fig05_agentic_workflow.png").exists():
        doc.add_picture(str(FIGURES_DIR / "fig05_agentic_workflow.png"), width=Inches(6.0))
        p_cap = doc.add_paragraph("Figure 6: Eight-Node LangGraph Autonomous Agent Execution Pipeline.")
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.runs[0].font.size = Pt(8.5)
        p_cap.runs[0].font.italic = True

    # References
    doc.add_heading("References", level=1)
    doc.add_paragraph(
        "[1] R. Anand and B. Singh, 'Effect of composition of board and promoter group retained ownership on underpricing of Indian IPO firms,' Indian Journal of Corporate Governance, vol. 12, no. 1, pp. 21–38, 2019.\n"
        "[2] B. Baba and G. Sevil, 'Predicting IPO initial returns using random forest,' Borsa Istanbul Review, vol. 20, no. 1, pp. 13–23, 2020.\n"
        "[3] E. Bajo and C. Raimondo, 'Media sentiment and IPO underpricing,' Journal of Corporate Finance, vol. 46, pp. 139–153, 2017.\n"
        "[4] R. Bansal and A. Khanna, 'Determinants of IPOs initial return: Extreme analysis of Indian market,' Journal of Financial Risk Management, vol. 1, no. 4, pp. 68–75, 2012.\n"
        "[5] E. Bastı, C. Kuzey, and D. Delen, 'Analyzing initial public offerings short-term performance using decision trees and SVMs,' Decision Support Systems, vol. 73, pp. 15–27, 2015.\n"
        "[6] T. Chen and C. Guestrin, 'XGBoost: A scalable tree boosting system,' in Proc. 22nd ACM SIGKDD Int. Conf. Knowledge Discovery and Data Mining, 2016, pp. 785–794.\n"
        "[7] S. Ghosh, 'Underpricing of initial public offerings: The Indian experience,' Emerging Markets Finance and Trade, vol. 41, no. 6, pp. 45–57, 2005.\n"
        "[8] S. Ghosh, A. Maji, N. H. Vardhan, and S. K. Naskar, 'Experimenting with multi-modal information to predict success of Indian IPOs,' arXiv preprint arXiv:2412.16174, 2024.\n"
        "[9] S. M. Lundberg and S.-I. Lee, 'A unified approach to interpreting model predictions,' Advances in Neural Information Processing Systems (NeurIPS), vol. 30, pp. 4765–4774, 2017."
    )

    doc.save(str(DOCX_PATH))
    print(f"DOCX successfully compiled to {DOCX_PATH}")


# =====================================================================
# 2. BUILD PDF
# =====================================================================
class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        canvas.Canvas.__init__(self, *args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748b"))
        
        # Header
        self.drawString(54, 750, "Pre-IPO Subscription Demand Forecasting & Explainability Research")
        self.setStrokeColor(colors.HexColor("#cbd5e1"))
        self.setLineWidth(0.5)
        self.line(54, 742, 558, 742)
        
        # Footer
        self.line(54, 50, 558, 50)
        page_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(558, 38, page_text)
        self.drawString(54, 38, "Leakage-Aware Machine Learning & Agentic AI Platform (v1.0.0)")
        self.restoreState()


def build_pdf():
    print("Generating PDF document (.pdf)...")
    doc = SimpleDocTemplate(
        str(PDF_PATH),
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        'PaperTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        textColor=colors.HexColor("#0f172a"),
        alignment=1, # Center
        spaceAfter=8
    )

    author_style = ParagraphStyle(
        'PaperAuthor',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=9,
        leading=12,
        textColor=colors.HexColor("#475569"),
        alignment=1,
        spaceAfter=14
    )

    h1_style = ParagraphStyle(
        'PaperH1',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=15,
        textColor=colors.HexColor("#0f172a"),
        spaceBefore=12,
        spaceAfter=6
    )

    body_style = ParagraphStyle(
        'PaperBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#1e293b"),
        spaceAfter=6
    )

    caption_style = ParagraphStyle(
        'PaperCaption',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#475569"),
        alignment=1,
        spaceAfter=10
    )

    story = []

    # Title & Authors
    story.append(Paragraph("Leakage-Aware Machine Learning and Agentic AI for Pre-IPO Subscription Demand Forecasting Across QIB, NII, Retail, and Total Subscription", title_style))
    story.append(Paragraph("Quantitative Financial Engineering & Computational Intelligence Group<br/>Dataset Version 1.1.0 • Model Registry 1.0.0-final-candidates", author_style))

    # Abstract Box Table
    abs_text = (
        "<b>Abstract— </b>Accurately forecasting subscription demand for Initial Public Offerings (IPOs) prior to bidding is a fundamental challenge in financial econometrics. "
        "We propose an ex-ante, leakage-aware Machine Learning (ML) and Agentic Artificial Intelligence (AI) framework to forecast category-wise subscription demand multiples "
        "across Qualified Institutional Buyers (QIB), Non-Institutional Investors (NII/HNI), Retail Individual Investors (RII), and Total demand for Indian Mainboard IPOs. "
        "Governed by a strict temporal barrier (t_observed &le; 09:59:59 IST on opening date), 45 features across 7 ablation groups were modeled using expanding walk-forward cross-validation. "
        "On a quarantined 2024 test set (N=17), the frozen Total model achieves <b>R² = +0.666</b> and <b>MAE = 17.43x</b>, while QIB achieves <b>R² = +0.589</b>. "
        "Incorporating pre-open GMP and media sentiment yields a <b>41.568% error reduction</b> over purely fundamental baselines. "
        "The system is encapsulated within an 8-node LangGraph orchestration graph, TreeSHAP explainability, and a React/TypeScript dashboard (98/98 tests passed).<br/><br/>"
        "<b>Keywords: </b><i>Initial Public Offerings (IPOs), Subscription Demand Forecasting, Temporal Data Leakage, TreeSHAP, Agentic AI, LangGraph, Decision Support.</i>"
    )
    abs_table = Table([[Paragraph(abs_text, body_style)]], colWidths=[500])
    abs_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#f1f5f9")),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#cbd5e1")),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('LEFTPADDING', (0, 0), (-1, -1), 10),
        ('RIGHTPADDING', (0, 0), (-1, -1), 10),
    ]))
    story.append(abs_table)
    story.append(Spacer(1, 10))

    # 1. Introduction
    story.append(Paragraph("1. Introduction & Research Problem", h1_style))
    story.append(Paragraph(
        "Initial Public Offerings (IPOs) represent a critical corporate milestone for capital formation. Under SEBI regulations in India, book-building issues allocate shares across Qualified Institutional Buyers (QIB), Non-Institutional Investors (NII), and Retail Individual Investors (RII). While prior research predominantly investigates post-listing underpricing, predicting pre-IPO subscription demand multiples across investor classes provides foundational intelligence for pricing, underwriting, and risk management.",
        body_style
    ))

    if (FIGURES_DIR / "fig01_methodology.png").exists():
        story.append(Image(str(FIGURES_DIR / "fig01_methodology.png"), width=480, height=280))
        story.append(Paragraph("Figure 1: Overall Pre-IPO Demand Forecasting & Explainability Research Pipeline.", caption_style))

    # 2. Temporal Anti-Leakage
    story.append(Paragraph("2. Strict Pre-IPO Temporal Anti-Leakage Framework", h1_style))
    story.append(Paragraph(
        "To eliminate temporal look-ahead contamination, all features must satisfy t_observed &le; 09:59:59 IST on the IPO opening date. All post-open order book tallies, intraday bids, and listing returns are strictly rejected.",
        body_style
    ))

    if (FIGURES_DIR / "fig02_temporal_leakage.png").exists():
        story.append(Image(str(FIGURES_DIR / "fig02_temporal_leakage.png"), width=480, height=210))
        story.append(Paragraph("Figure 2: Strict Pre-IPO Ex-Ante Temporal Anti-Leakage Boundary.", caption_style))

    # 3. Model Results Table
    story.append(Paragraph("3. Quarantined Test Evaluation & Model Performance", h1_style))
    table_data = [
        [Paragraph("<b>Target Category</b>", body_style), Paragraph("<b>Selected Model</b>", body_style), Paragraph("<b>Transform</b>", body_style), Paragraph("<b>Test MAE</b>", body_style), Paragraph("<b>Test RMSE</b>", body_style), Paragraph("<b>Test R²</b>", body_style)],
        ["Total Subscription", "XGBoost Regressor", "log1p", "17.43x", "22.92x", "+0.666"],
        ["QIB Subscription", "Random Forest", "RAW", "31.43x", "38.48x", "+0.589"],
        ["NII Subscription", "LightGBM Regressor", "log1p", "39.62x", "64.44x", "+0.385"],
        ["Retail Subscription", "XGBoost Regressor", "RAW", "23.14x", "36.71x", "-0.450"],
    ]
    t_res = Table(table_data, colWidths=[100, 110, 60, 70, 70, 70])
    t_res.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0f172a")),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ALIGN', (3, 1), (-1, -1), 'CENTER'),
    ]))
    story.append(t_res)
    story.append(Spacer(1, 10))

    if (FIGURES_DIR / "fig08_actual_vs_predicted.png").exists():
        story.append(Image(str(FIGURES_DIR / "fig08_actual_vs_predicted.png"), width=340, height=290))
        story.append(Paragraph("Figure 3: Actual vs. Predicted Total Subscription (R² = +0.666, N=17).", caption_style))

    # 4. TreeSHAP & LangGraph
    story.append(Paragraph("4. TreeSHAP Feature Attribution & Agentic AI Architecture", h1_style))
    story.append(Paragraph(
        "TreeSHAP additive decompositions reveal that forward market indicators (GMP, positive news sentiment) dominate institutional and aggregate predictions, while offering scale friction dominates retail predictions. The quantitative core is orchestrated via an 8-node LangGraph StateGraph delivering real-time decision support.",
        body_style
    ))

    if (FIGURES_DIR / "fig07_shap_importance.png").exists():
        story.append(Image(str(FIGURES_DIR / "fig07_shap_importance.png"), width=440, height=240))
        story.append(Paragraph("Figure 4: Global TreeSHAP Feature Attributions for Total Subscription.", caption_style))

    if (FIGURES_DIR / "fig05_agentic_workflow.png").exists():
        story.append(Image(str(FIGURES_DIR / "fig05_agentic_workflow.png"), width=480, height=180))
        story.append(Paragraph("Figure 5: Eight-Node LangGraph Autonomous Agent Execution Pipeline.", caption_style))

    # References
    story.append(Paragraph("References", h1_style))
    ref_text = (
        "[1] R. Anand and B. Singh, 'Effect of board composition on IPO underpricing,' Indian J. Corp. Gov., 2019.<br/>"
        "[2] B. Baba and G. Sevil, 'Predicting IPO initial returns using random forest,' Borsa Istanbul Rev., 2020.<br/>"
        "[3] E. Bajo and C. Raimondo, 'Media sentiment and IPO underpricing,' J. Corp. Finance, 2017.<br/>"
        "[4] T. Chen and C. Guestrin, 'XGBoost: A scalable tree boosting system,' ACM SIGKDD, 2016.<br/>"
        "[5] S. Ghosh, 'Underpricing of initial public offerings: The Indian experience,' Emerg. Mark. Finance Trade, 2005.<br/>"
        "[6] S. Ghosh et al., 'Experimenting with multi-modal information to predict success of Indian IPOs,' arXiv:2412.16174, 2024.<br/>"
        "[7] S. M. Lundberg and S.-I. Lee, 'A unified approach to interpreting model predictions,' NeurIPS, 2017."
    )
    story.append(Paragraph(ref_text, body_style))

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"PDF successfully compiled to {PDF_PATH}")


if __name__ == "__main__":
    build_docx()
    build_pdf()
